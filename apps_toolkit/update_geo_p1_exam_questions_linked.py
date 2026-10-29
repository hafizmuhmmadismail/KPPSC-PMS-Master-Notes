import os
import re
import sys
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
WORDCRAFT_CLI = r"H:\My Drive\PMS KPK 2022 notes\apps_toolkit\tools\wordcraft\wordcraft-0.8.0-windows-x64-portable\wordcraft-cli.exe"
GEO_P1_DIR = r"H:\My Drive\PMS KPK 2022 notes\3. Geography Paper 1 and Paper 2\Paper 1 - Physical and General Geography"

GEO_P1_EXAM_QUESTIONS_LINKED_MAP = {
    1: (
        "• 2025 Q.1: \"Describe the internal structure of the Earth and explain how seismic wave behavior provides evidence for the layering of the Earth's interior. Evaluate the scientific reliability of seismic methods in understanding deep Earth processes. (20 Marks)\"\n"
        "• 2016 Q.3: \"What are the different layers of the Earth? Relate the three types of plate collisions associated with seismicity, volcanism, and mountain building. (20 Marks)\"\n"
        "• 2008 Q.5: \"Define the internal structure of the Earth from Asthenosphere to inner core. Give reasons why the core is composed of iron. (20 Marks)\"\n"
        "• 2008 Q.2: \"Describe the mechanism and types of seismic waves. Define the causes of earthquakes and write an account of earthquakes in Pakistan. (20 Marks)\""
    ),
    2: (
        "• 2016 Q.3: \"What are the different layers of the Earth? Relate the three types of plate collisions associated with seismicity, volcanism, and mountain building. (20 Marks)\"\n"
        "• 2010 Q.4: \"What is an earthquake, and how does it arise? Describe the relationship between earthquakes and plate tectonics. (20 Marks)\"\n"
        "• 2008 Q.2: \"Describe the mechanism and types of seismic waves. Define the causes of earthquakes and write an account of earthquakes in Pakistan. (20 Marks)\""
    ),
    3: (
        "• 2022 Q.2: \"Provide a detailed account of sedimentary rocks and their most workable classification based on origin and composition. (20 Marks)\"\n"
        "• 2018 Q.3: \"Discuss the main types of rocks along with examples and describe the geological rock cycle. (20 Marks)\"\n"
        "• 2013 Q.3: \"Define rocks and minerals. Discuss the circumstances under which sedimentary rocks are formed and classified. (20 Marks)\"\n"
        "• 2008 Q.1: \"Describe the rock cycle. How can we identify igneous rocks on the basis of mineral contents? Explain types of sedimentary rocks. (20 Marks)\""
    ),
    4: (
        "• 2022 Q.1: \"Discuss various processes of physical and chemical weathering with support from sketches/diagrams. Explain their geomorphic significance. (20 Marks)\"\n"
        "• 2018 Q.1: \"Discuss the endogenic and exogenic geomorphic processes in detail. How do they shape the Earth's relief? (20 Marks)\"\n"
        "• 2013 Q.4: \"Discuss the factors promoting mass movements in weathered materials. Explain the various types of mass movements. (20 Marks)\""
    ),
    5: (
        "• 2018 Q.2: \"What are the major erosional and depositional landforms produced by rivers? Illustrate their evolution from youth to old age. (20 Marks)\"\n"
        "• 2008 Q.7: \"Explain the stages of river valley development and describe the erosional and depositional landforms associated with running water. (20 Marks)\""
    ),
    6: (
        "• 2022 Q.6: \"Discuss valley glaciers and the resultant erosional and depositional landforms in detail. (20 Marks)\"\n"
        "• 2018 Q.8(a): \"Write short note on: Glacial landforms. (10 Marks)\"\n"
        "• 2010 Q.3: \"Differentiate between Valley Glacier and Ice Sheet. Discuss the depositional landforms of Valley Glacier in detail. (20 Marks)\""
    ),
    7: (
        "• 2025 Q.6: \"Discuss the environmental conditions necessary for the development of Karst Landscapes. Describe the major Karst Landforms and evaluate the processes responsible for their evolution. (20 Marks)\"\n"
        "• 2016 Q.4: \"Describe the nature of groundwater and define the elements of groundwater environment. Review the processes and features associated with Karst topography. (20 Marks)\"\n"
        "• 2010 Q.8(a): \"Write short note on: Karst Topography. (10 Marks)\""
    ),
    8: (
        "• 2016 Q.6: \"\\\"The coastal system is the scene of complex tidal fluctuation, winds, waves, and ocean currents.\\\" Discuss coastal processes and coastal landforms. (20 Marks)\"\n"
        "• 2016 Q.8(c): \"Write short note on: Wind deposition / Aeolian landforms. (10 Marks)\""
    ),
    9: (
        "• 2016 Q.1: \"\\\"The heat energy entering the atmosphere and its part reaching the earth surface is balanced by an equal amount of heat going out into outer space.\\\" Discuss the global heat balance. (20 Marks)\"\n"
        "• 2013 Q.1: \"Write an essay on the Radiation Balance of the Earth and Atmosphere. (20 Marks)\"\n"
        "• 2010 Q.1: \"Write an essay on the influence of latitude on the annual cycle of insolation. (20 Marks)\"\n"
        "• 2008 Q.3: \"\\\"Solar input energy is equal to Earth's output energy.\\\" Explain. Describe the classification of air masses with examples. (20 Marks)\""
    ),
    10: (
        "• 2022 Q.3: \"Discuss the distribution of pressure belts over the globe and the resulting planetary wind circulation systems. (20 Marks)\"\n"
        "• 2018 Q.5: \"Describe the different layers and composition of the atmosphere with reference to lapse rate and pause zones. (20 Marks)\"\n"
        "• 2016 Q.8(a): \"Write short note on: Planetary wind system. (10 Marks)\"\n"
        "• 2013 Q.2: \"What are the causes of Atmospheric Circulation? Explain the large-scale wind systems of the globe. (20 Marks)\""
    ),
    11: (
        "• 2025 Q.3: \"Discuss the major forms of precipitation and the atmospheric mechanisms that produce them. Support your answer with suitable regional and local examples illustrating each type. (20 Marks)\"\n"
        "• 2018 Q.4: \"Discuss the horizontal distribution of temperature and factors affecting the temperature distribution across the globe. (20 Marks)\""
    ),
    12: (
        "• 2022 Q.4: \"Air masses and fronts: Give a comprehensive classification and elaborate their weather characteristics with sketches. (20 Marks)\"\n"
        "• 2016 Q.2: \"What are the distinguishing characteristics of an air mass? How does a source region influence the type of an air mass that forms over it? Explain the life cycle of a wave cyclone. (20 Marks)\"\n"
        "• 2010 Q.2: \"Define Cyclone and Anticyclone with particular reference to air movement and circulation in Northern and Southern hemispheres. What weather is associated with each? (20 Marks)\"\n"
        "• 2008 Q.3: \"\\\"Solar input energy is equal to Earth's output energy.\\\" Explain. Describe the classification of air masses with examples. (20 Marks)\""
    ),
    13: (
        "• 2025 Q.4: \"Analyze the formation and development of tropical weather systems with reference to hurricanes. Describe the weather patterns associated with hurricanes and the environmental factors influencing their intensity. (20 Marks)\"\n"
        "• 2010 Q.2: \"Define Cyclone and Anticyclone with particular reference to air movement and circulation in Northern and Southern hemispheres. What weather is associated with each? (20 Marks)\""
    ),
    14: (
        "• 2025 Q.5: \"Explain the basis of the First Letter Classification in the Köppen Climate Classification System. Compare and contrast the Af (tropical rainforest) and Aw (tropical savanna) climates in terms of distribution, temperature-precipitation regimes, and ecological characteristics. (20 Marks)\"\n"
        "• 2016 Q.8(b): \"Write short note on: The Köppen climatic classification. (10 Marks)\""
    ),
    15: (
        "• 2016 Q.5: \"Explain the chemical composition of seawater and the basic physical structure of the ocean. (20 Marks)\"\n"
        "• 2010 Q.8(b): \"Write short note on: Ocean floor morphology. (10 Marks)\""
    ),
    16: (
        "• 2025 Q.8: \"Describe the factors controlling salinity of sea water. Critically analyze the spatial and vertical variations in salinity across the world's oceans and explain their influence on ocean circulation and marine life. (20 Marks)\"\n"
        "• 2022 Q.7: \"Analyze the major causes of salinity variation in the oceans of the world. (20 Marks)\"\n"
        "• 2013 Q.6: \"Write a comprehensive note on the salinity of ocean water and factors governing its variations. (20 Marks)\"\n"
        "• 2010 Q.6: \"Why does salinity of ocean water vary from place to place? What are the determining factors? (20 Marks)\""
    ),
    17: (
        "• 2022 Q.8(d): \"Write short note on: Ocean currents of the Atlantic Ocean. (10 Marks)\"\n"
        "• 2018 Q.6: \"Explain with the help of maps the major ocean currents of the Indian, Pacific, and Atlantic Oceans. (20 Marks)\"\n"
        "• 2013 Q.5: \"Describe the currents of the Pacific Ocean and their climatic significance. (20 Marks)\"\n"
        "• 2008 Q.4: \"Explain the mechanism of oceanic currents that produce circulation in ocean water. (20 Marks)\""
    ),
    18: (
        "• 2025 Q.2: \"Explain the origin of ocean tides and critically examine the major theories proposed to explain tidal behavior. How do local coastal configurations influence tidal patterns? (20 Marks)\"\n"
        "• 2018 Q.8(c): \"Write short note on: Ocean tides. (10 Marks)\"\n"
        "• 2016 Q.6: \"\\\"The coastal system is the scene of complex tidal fluctuation, winds, waves, and ocean currents.\\\" Discuss coastal processes and coastal landforms. (20 Marks)\""
    ),
    19: (
        "• 2025 Q.7: \"Explain the principles behind the construction and interpretation of weather maps. How are modern weather forecasts generated, and what are the limitations of current forecasting techniques? (20 Marks)\"\n"
        "• 2018 Q.7: \"Define map and discuss its types along with their practical uses in modern geographical studies. (20 Marks)\"\n"
        "• 2016 Q.7: \"Define cartography and mapping basics: Explain map scale (Representative Fraction) and the principles of map projections. (20 Marks)\"\n"
        "• 2010 Q.7: \"Define Map. What are the essentials of a map? Why do we need projections for drawing a map? (20 Marks)\""
    ),
    20: (
        "• 2022 Q.8(b/c): \"Write comprehensive short notes on: Merits and demerits of cylindrical projections / Map projections and their choice in cartography. (10 + 10 = 20 Marks)\"\n"
        "• 2018 Q.8(b): \"Write short note on: Mercator’s Projection. (10 Marks)\"\n"
        "• 2016 Q.8(d): \"Write short note on: Mercator’s projection. (10 Marks)\"\n"
        "• 2013 Q.7: \"What are Map Projections? Give the merits and demerits of Cylindrical Projections. (20 Marks)\"\n"
        "• 2010 Q.8(d): \"Write short note on: Mercator's Projection. (10 Marks)\"\n"
        "• 2008 Q.6: \"How can we identify different types of maps? Characteristics of Cylindrical and Conical Projections. (20 Marks)\""
    )
}

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text_list, title="EXAM ESSENTIAL TAKEAWAY"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"\U0001F4CC {title}\n")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10.5)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(14, 43, 80)
    
    for line in text_list:
        p_line = cell.add_paragraph()
        p_line.paragraph_format.space_before = Pt(2)
        p_line.paragraph_format.space_after = Pt(2)
        run = p_line.add_run(line)
        run.font.name = "Georgia"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 41, 59)

def generate_geo_master_doc(meta_data, callout_data, map_text, timeline_data, essay_sections, quotes_data, out_path, paper_sub_title):
    doc = Document()
    
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_main = p_title.add_run("KPPSC PMS • GEOGRAPHY (PAPER-I)")
    run_main.font.name = "Arial"
    run_main.font.size = Pt(12)
    run_main.font.bold = True
    run_main.font.color.rgb = RGBColor(14, 43, 80)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(10)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run(f"EXECUTIVE MASTER EXAMINATION NOTES • {paper_sub_title}")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(27, 73, 101)
    
    # Metadata Table
    tbl_meta = doc.add_table(rows=len(meta_data), cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_meta.autofit = False
    
    for idx, (k, v) in enumerate(meta_data):
        row = tbl_meta.rows[idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        
        set_cell_margins(cell_k, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_v, top=60, bottom=60, left=100, right=100)
        
        p_k = cell_k.paragraphs[0]
        p_k.paragraph_format.space_before = Pt(0)
        p_k.paragraph_format.space_after = Pt(0)
        r_k = p_k.add_run(k)
        r_k.font.name = "Arial"
        r_k.font.size = Pt(9)
        r_k.font.bold = True
        
        lines = v.split("\n")
        for l_idx, line in enumerate(lines):
            if l_idx == 0:
                p_v = cell_v.paragraphs[0]
            else:
                p_v = cell_v.add_paragraph()
            p_v.paragraph_format.space_before = Pt(1 if len(lines) > 2 else 0)
            p_v.paragraph_format.space_after = Pt(1 if len(lines) > 2 else 0)
            r_v = p_v.add_run(line)
            r_v.font.name = "Georgia"
            r_v.font.size = Pt(8.5 if len(lines) > 2 else 9)
            if idx == 0:
                r_v.font.color.rgb = RGBColor(255, 255, 255)
                r_v.font.bold = True
            else:
                r_v.font.color.rgb = RGBColor(15, 23, 42)
                
        if idx == 0:
            set_cell_background(cell_k, "0E2B50")
            set_cell_background(cell_v, "0E2B50")
            r_k.font.color.rgb = RGBColor(255, 255, 255)
        else:
            bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
            set_cell_background(cell_k, bg)
            set_cell_background(cell_v, bg)
            r_k.font.color.rgb = RGBColor(14, 43, 80) if "EXAM QUESTIONS" in k else RGBColor(30, 41, 59)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_callout(doc, callout_data, "EXAM ESSENTIAL TAKEAWAY")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # 1. Conceptual Mind Map
    p_hm = doc.add_paragraph()
    p_hm.paragraph_format.space_before = Pt(8)
    p_hm.paragraph_format.space_after = Pt(4)
    r_hm = p_hm.add_run("1. CONCEPTUAL MIND MAP & THEMATIC ROADMAP")
    r_hm.font.name = "Arial"
    r_hm.font.size = Pt(11)
    r_hm.font.bold = True
    r_hm.font.color.rgb = RGBColor(14, 43, 80)
    
    tbl_mm = doc.add_table(rows=1, cols=1)
    tbl_mm.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mm.autofit = False
    c_mm = tbl_mm.cell(0, 0)
    c_mm.width = Inches(6.5)
    set_cell_background(c_mm, "F8FAFC")
    set_cell_margins(c_mm, top=100, bottom=100, left=140, right=140)
    p_mm_c = c_mm.paragraphs[0]
    p_mm_c.paragraph_format.space_before = Pt(0)
    p_mm_c.paragraph_format.space_after = Pt(0)
    r_mm_code = p_mm_c.add_run(map_text)
    r_mm_code.font.name = "Consolas"
    r_mm_code.font.size = Pt(8.5)
    r_mm_code.font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # 2. Analytical Dimensions Matrix
    p_htl = doc.add_paragraph()
    p_htl.paragraph_format.space_before = Pt(8)
    p_htl.paragraph_format.space_after = Pt(4)
    r_htl = p_htl.add_run("2. ANALYTICAL DIMENSIONS & GEOMORPHIC/SYSTEMIC MATRIX")
    r_htl.font.name = "Arial"
    r_htl.font.size = Pt(11)
    r_htl.font.bold = True
    r_htl.font.color.rgb = RGBColor(14, 43, 80)
    
    col_count = len(timeline_data[0]) if timeline_data else 3
    tbl_tl = doc.add_table(rows=len(timeline_data), cols=col_count)
    tbl_tl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_tl.autofit = False
    
    if col_count == 4:
        widths = [Inches(1.5), Inches(1.8), Inches(2.0), Inches(1.2)]
    else:
        widths = [Inches(1.8), Inches(2.3), Inches(2.4)]
        
    for r_idx, row_data in enumerate(timeline_data):
        row = tbl_tl.rows[r_idx]
        is_hdr = (r_idx == 0)
        bg = "0E2B50" if is_hdr else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
        fg = RGBColor(255, 255, 255) if is_hdr else RGBColor(30, 41, 59)
        
        for c_idx, text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Arial" if is_hdr else "Georgia"
            r.font.size = Pt(8.5 if not is_hdr else 9)
            r.font.bold = is_hdr or (c_idx == 0)
            r.font.color.rgb = fg
            
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # 3. Model Essay Section
    p_he = doc.add_paragraph()
    p_he.paragraph_format.space_before = Pt(12)
    p_he.paragraph_format.space_after = Pt(4)
    r_he = p_he.add_run("3. THE 20-MARK MASTER MODEL ESSAY (~650-700 WORDS)")
    r_he.font.name = "Arial"
    r_he.font.size = Pt(11)
    r_he.font.bold = True
    r_he.font.color.rgb = RGBColor(14, 43, 80)
    
    for sub_h, body_p in essay_sections:
        p_sh = doc.add_paragraph()
        p_sh.paragraph_format.space_before = Pt(8)
        p_sh.paragraph_format.space_after = Pt(2)
        r_sh = p_sh.add_run(sub_h)
        r_sh.font.name = "Arial"
        r_sh.font.size = Pt(10)
        r_sh.font.bold = True
        r_sh.font.color.rgb = RGBColor(27, 73, 101)
        
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(4)
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_b = p_b.add_run(body_p)
        r_b.font.name = "Georgia"
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = RGBColor(30, 41, 59)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # 4. Scholarly Perspectives Table
    p_hq = doc.add_paragraph()
    p_hq.paragraph_format.space_before = Pt(12)
    p_hq.paragraph_format.space_after = Pt(4)
    r_hq = p_hq.add_run("4. SCHOLARLY PERSPECTIVES & GEOMORPHIC AXIOMS")
    r_hq.font.name = "Arial"
    r_hq.font.size = Pt(11)
    r_hq.font.bold = True
    r_hq.font.color.rgb = RGBColor(14, 43, 80)
    
    q_widths = [Inches(1.8), Inches(1.8), Inches(2.9)]
    tbl_qt = doc.add_table(rows=len(quotes_data), cols=3)
    tbl_qt.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_qt.autofit = False
    
    for r_idx, row_data in enumerate(quotes_data):
        row = tbl_qt.rows[r_idx]
        is_hdr = (r_idx == 0)
        bg = "0E2B50" if is_hdr else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
        fg = RGBColor(255, 255, 255) if is_hdr else RGBColor(30, 41, 59)
        
        for c_idx, text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = q_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Arial" if is_hdr else "Georgia"
            r.font.size = Pt(8.5 if not is_hdr else 9)
            r.font.bold = is_hdr or (c_idx == 0)
            if not is_hdr and c_idx == 2:
                r.font.italic = True
            r.font.color.rgb = fg
            
    doc.save(out_path)
    print(f"Generated Master Docx: {out_path}")

def generate_pdf_and_md(meta_data, callout_data, map_text, timeline_data, essay_sections, quotes_data, pdf_path, md_path, title):
    md_content = f"# {title}\n\n"
    md_content += "## KPPSC PMS Examination — Master Study Notes (Paper-I: Physical & General Geography)\n\n"
    
    md_content += "| Field | Value |\n|---|---|\n"
    for k, v in meta_data:
        v_table = "<br>".join(v.split("\n"))
        md_content += f"| **{k}** | {v_table} |\n"
    md_content += "\n---\n\n"
    
    # Dedicated EXAM QUESTIONS LINKED Section in MD
    exam_questions = []
    for k, v in meta_data:
        if "EXAM QUESTIONS" in k.upper():
            exam_questions = v.split("\n")
            break
            
    if exam_questions:
        md_content += "### 🎯 EXAM QUESTIONS LINKED\n\n"
        for q in exam_questions:
            md_content += f"> {q}\n>\n"
        md_content += "\n---\n\n"
        
    md_content += "### 📌 EXAM ESSENTIAL TAKEAWAY\n"
    for line in callout_data:
        md_content += f"> {line}\n"
    md_content += "\n---\n\n"
    
    md_content += "### 1. Conceptual Mind Map & Thematic Roadmap\n```\n" + map_text + "\n```\n\n---\n\n"
    
    col_count = len(timeline_data[0]) if timeline_data else 3
    md_content += "### 2. Analytical Dimensions & Geomorphic/Systemic Matrix\n\n"
    headers = timeline_data[0]
    if col_count == 4:
        md_content += f"| {headers[0]} | {headers[1]} | {headers[2]} | {headers[3]} |\n|---|---|---|---|\n"
        for row in timeline_data[1:]:
            md_content += f"| **{row[0]}** | {row[1]} | {row[2]} | {row[3]} |\n"
    else:
        md_content += f"| {headers[0]} | {headers[1]} | {headers[2]} |\n|---|---|---|\n"
        for row in timeline_data[1:]:
            md_content += f"| **{row[0]}** | {row[1]} | {row[2]} |\n"
    md_content += "\n---\n\n"
    
    md_content += "### 3. 20-Mark Master Model Essay\n\n"
    for sh, bd in essay_sections:
        md_content += f"#### {sh}\n{bd}\n\n"
    md_content += "---\n\n"
    
    md_content += "### 4. Scholarly Perspectives & Geomorphic Axioms\n\n"
    q_hdr = quotes_data[0]
    md_content += f"| {q_hdr[0]} | {q_hdr[1]} | {q_hdr[2]} |\n|---|---|---|\n"
    for row in quotes_data[1:]:
        md_content += f"| **{row[0]}** | {row[1]} | *\"{row[2]}\"* |\n"
        
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md_content)
        
    # Build HTML for PDF
    html_content = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
    @page {{ size: A4; margin: 18mm 16mm 18mm 16mm; }}
    body {{ font-family: 'Segoe UI', Calibri, Arial, sans-serif; font-size: 9.5pt; line-height: 1.45; color: #1A1A1A; margin: 0; }}
    h1 {{ color: #0E2B50; font-size: 15pt; border-bottom: 2px solid #0E2B50; padding-bottom: 3px; margin: 0 0 8px 0; text-align: center; }}
    h2 {{ color: #1B4965; font-size: 11pt; margin: 10px 0 6px 0; border-bottom: 1px solid #D0D7DE; padding-bottom: 2px; text-align: center; }}
    h3 {{ color: #0E2B50; font-size: 10.5pt; margin: 12px 0 4px 0; }}
    h4 {{ color: #1B4965; font-size: 10pt; margin: 8px 0 2px 0; }}
    table {{ width: 100%; border-collapse: collapse; margin: 8px 0; font-size: 8.5pt; }}
    th {{ background-color: #0E2B50; color: #FFFFFF; padding: 4px 8px; border: 1px solid #0E2B50; text-align: left; }}
    td {{ padding: 4px 8px; border: 1px solid #D0D7DE; vertical-align: top; }}
    tr:nth-child(even) td {{ background-color: #F8FAFC; }}
    blockquote {{ border-left: 4px solid #0E2B50; background-color: #F0F4F8; padding: 6px 12px; margin: 8px 0; font-size: 9pt; }}
    .callout-box {{ background-color: #F0F4F8; border-left: 4px solid #0E2B50; padding: 8px 12px; margin: 8px 0; font-size: 8.5pt; }}
    pre {{ background-color: #F8FAFC; border: 1px solid #E2E8F0; padding: 8px; font-family: 'Consolas', monospace; font-size: 8pt; line-height: 1.3; overflow-x: auto; }}
    p {{ margin: 4px 0 6px 0; text-align: justify; }}
    .footer {{ margin-top: 20px; font-size: 8pt; color: #666; text-align: center; border-top: 1px solid #E1E4E8; padding-top: 4px; }}
</style>
</head>
<body>
<h1>{title}</h1>
<h2>KPPSC PMS Master Examination Notes — Paper-I (Physical & General Geography)</h2>

<table>
<tbody>
"""
    for k, v in meta_data:
        v_html = "<br/>".join([f"<strong>{line[:line.find(':')+1]}</strong>{line[line.find(':')+1:]}" if ":" in line and any(yr in line for yr in ["2025", "2022", "2018", "2016", "2013", "2010", "2008"]) else line for line in v.split("\n")])
        html_content += f"<tr><th style='width: 30%;'>{k}</th><td>{v_html}</td></tr>\n"
        
    html_content += """</tbody>
</table>
"""
    if exam_questions:
        html_content += """<h2>🎯 Exam Questions Linked</h2>
<div class="callout-box">
"""
        for q in exam_questions:
            q_clean = q.strip()
            if ":" in q_clean:
                prefix, text = q_clean.split(":", 1)
                html_content += f"<p style='margin: 4px 0;'><strong>{prefix}:</strong>{text}</p>\n"
            else:
                html_content += f"<p style='margin: 4px 0;'>{q_clean}</p>\n"
        html_content += "</div>\n"

    html_content += """<blockquote>
"""
    for line in callout_data:
        html_content += f"<p>{line}</p>\n"
    html_content += f"""</blockquote>

<h3>1. Conceptual Mind Map & Thematic Roadmap</h3>
<pre>{map_text}</pre>

<h3>2. Analytical Dimensions & Geomorphic/Systemic Matrix</h3>
<table>
<thead><tr>
"""
    for h in timeline_data[0]:
        html_content += f"<th>{h}</th>"
    html_content += """</tr></thead>
<tbody>
"""
    for row in timeline_data[1:]:
        html_content += "<tr>"
        for idx, val in enumerate(row):
            if idx == 0:
                html_content += f"<td><strong>{val}</strong></td>"
            else:
                html_content += f"<td>{val}</td>"
        html_content += "</tr>\n"
    html_content += f"""</tbody>
</table>

<h3>3. 20-Mark Master Model Essay</h3>
"""
    for sh, bd in essay_sections:
        html_content += f"<h4>{sh}</h4>\n<p>{bd}</p>\n"

    html_content += f"""<h3>4. Scholarly Perspectives & Geomorphic Axioms</h3>
<table>
<thead><tr>
"""
    for h in quotes_data[0]:
        html_content += f"<th>{h}</th>"
    html_content += f"""</tr></thead>
<tbody>
"""
    for row in quotes_data[1:]:
        html_content += f"<tr><td><strong>{row[0]}</strong></td><td>{row[1]}</td><td><em>\"{row[2]}\"</em></td></tr>\n"
    html_content += f"""</tbody>
</table>

<div class="footer">
Khyber Pakhtunkhwa Public Service Commission (KPPSC) • PMS Geography Paper-I Master Notes (2025 Edition)
</div>
</body>
</html>"""

    temp_html = md_path.replace(".md", "_temp.html")
    with open(temp_html, "w", encoding="utf-8") as f:
        f.write(html_content)
        
    cmd = [EDGE_PATH, "--headless", "--disable-gpu", "--run-all-compositor-stages-before-draw", f"--print-to-pdf={pdf_path}", temp_html]
    subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    if os.path.exists(temp_html):
        os.remove(temp_html)
    print(f"PDF and MD successfully generated: {pdf_path}")


def extract_data_from_file(py_path):
    with open(py_path, 'r', encoding='utf-8') as f:
        code = f.read()
    
    clean_lines = []
    for line in code.splitlines():
        if re.match(r'^(generate_geo_master_doc|generate_pdf_and_md|for item in batch)', line):
            break
        clean_lines.append(line)
    clean_code = "\n".join(clean_lines)
    
    ns = {}
    exec(clean_code, ns)
    return ns

def update_and_build_all():
    script_map = [
        (r"H:\My Drive\PMS KPK 2022 notes\apps_toolkit\build_geo_p1_batch1_notes.py", [1, 2, 3, 4, 5, 6]),
        (r"H:\My Drive\PMS KPK 2022 notes\apps_toolkit\build_geo_p1_batch2_notes.py", [7, 8, 9, 10, 11, 12, 13, 14]),
        (r"H:\My Drive\PMS KPK 2022 notes\apps_toolkit\build_geo_p1_batch3_notes.py", [15, 16, 17, 18, 19, 20]),
    ]
    
    titles_map = {
        1: ("Q1_Internal_Structure_of_Earth_and_Seismic_Evidence_Master_Notes", "QUESTION 1: INTERNAL STRUCTURE OF EARTH & SEISMIC EVIDENCE", "KPPSC PMS Geography Paper-I — Q.1: Internal Structure of Earth & Seismic Discontinuities"),
        2: ("Q2_Plate_Tectonics_Theory_and_Orogenesis_Master_Notes", "QUESTION 2: PLATE TECTONICS THEORY & OROGENESIS", "KPPSC PMS Geography Paper-I — Q.2: Plate Tectonics Theory, Plate Collisions & Orogenesis"),
        3: ("Q3_Rocks_Classification_and_Rock_Cycle_Master_Notes", "QUESTION 3: ROCKS CLASSIFICATION & THE ROCK CYCLE", "KPPSC PMS Geography Paper-I — Q.3: Rocks, Sedimentary Classification & The Rock Cycle"),
        4: ("Q4_Weathering_Processes_and_Mass_Wasting_Master_Notes", "QUESTION 4: WEATHERING PROCESSES & MASS WASTING", "KPPSC PMS Geography Paper-I — Q.4: Weathering Processes (Physical & Chemical) & Mass Wasting"),
        5: ("Q5_Fluvial_Geomorphology_and_River_Landforms_Master_Notes", "QUESTION 5: FLUVIAL GEOMORPHOLOGY & RIVER LANDFORMS", "KPPSC PMS Geography Paper-I — Q.5: Fluvial Geomorphology & River Valley Evolution"),
        6: ("Q6_Glacial_Geomorphology_Valley_Glaciers_vs_Ice_Sheets_Master_Notes", "QUESTION 6: GLACIAL GEOMORPHOLOGY: VALLEY GLACIERS VS ICE SHEETS", "KPPSC PMS Geography Paper-I — Q.6: Glacial Geomorphology: Valley Glaciers vs. Ice Sheets"),
        7: ("Q7_Karst_Topography_and_Groundwater_Hydrology_Master_Notes", "QUESTION 7: KARST TOPOGRAPHY & GROUNDWATER HYDROLOGY", "KPPSC PMS Geography Paper-I — Q.7: Karst Topography & Groundwater Hydrology"),
        8: ("Q8_Aeolian_and_Coastal_Geomorphology_Master_Notes", "QUESTION 8: AEOLIAN & COASTAL GEOMORPHOLOGY", "KPPSC PMS Geography Paper-I — Q.8: Aeolian & Coastal Geomorphology"),
        9: ("Q9_Solar_Insolation_Heat_Budget_and_Radiation_Balance_Master_Notes", "QUESTION 9: SOLAR INSOLATION, HEAT BUDGET & RADIATION BALANCE", "KPPSC PMS Geography Paper-I — Q.9: Solar Insolation, Heat Budget & Radiation Balance"),
        10: ("Q10_Atmospheric_Pressure_Belts_and_Planetary_Winds_Master_Notes", "QUESTION 10: ATMOSPHERIC PRESSURE BELTS & PLANETARY WINDS", "KPPSC PMS Geography Paper-I — Q.10: Atmospheric Pressure Belts & Planetary Winds"),
        11: ("Q11_Atmospheric_Moisture_Clouds_and_Precipitation_Master_Notes", "QUESTION 11: ATMOSPHERIC MOISTURE, CLOUDS & PRECIPITATION", "KPPSC PMS Geography Paper-I — Q.11: Atmospheric Moisture, Clouds & Precipitation"),
        12: ("Q12_Air_Masses_Frontogenesis_and_Wave_Cyclones_Master_Notes", "QUESTION 12: AIR MASSES, FRONTOGENESIS & POLAR FRONT CYCLONES", "KPPSC PMS Geography Paper-I — Q.12: Air Masses, Frontogenesis & Polar Front Wave Cyclones"),
        13: ("Q13_Tropical_Weather_Systems_and_Hurricane_Dynamics_Master_Notes", "QUESTION 13: TROPICAL WEATHER SYSTEMS & HURRICANE DYNAMICS", "KPPSC PMS Geography Paper-I — Q.13: Tropical Weather Systems & Hurricane Dynamics"),
        14: ("Q14_Koppen_Climatic_Classification_Af_vs_Aw_Master_Notes", "QUESTION 14: KÖPPEN CLIMATIC CLASSIFICATION & Af VS Aw", "KPPSC PMS Geography Paper-I — Q.14: Köppen Climatic Classification & Af vs. Aw Climates"),
        15: ("Q15_Ocean_Floor_Morphology_and_Continental_Margins_Master_Notes", "QUESTION 15: OCEAN FLOOR MORPHOLOGY & CONTINENTAL MARGINS", "KPPSC PMS Geography Paper-I — Q.15: Ocean Floor Morphology & Continental Margins"),
        16: ("Q16_Ocean_Salinity_Controlling_Factors_and_Variations_Master_Notes", "QUESTION 16: OCEAN SALINITY: CONTROLLING FACTORS & VARIATIONS", "KPPSC PMS Geography Paper-I — Q.16: Ocean Salinity: Factors, Variations & Circulation"),
        17: ("Q17_Oceanic_Circulation_Major_Currents_and_Gyres_Master_Notes", "QUESTION 17: OCEANIC CIRCULATION: MAJOR CURRENTS & GYRES", "KPPSC PMS Geography Paper-I — Q.17: Oceanic Circulation: Major Ocean Currents & Gyres"),
        18: ("Q18_Mechanics_of_Ocean_Tides_and_Tidal_Theories_Master_Notes", "QUESTION 18: MECHANICS OF OCEAN TIDES & TIDAL THEORIES", "KPPSC PMS Geography Paper-I — Q.18: Mechanics of Ocean Tides & Tidal Theories"),
        19: ("Q19_Cartography_Basics_Map_Scales_and_Weather_Maps_Master_Notes", "QUESTION 19: CARTOGRAPHY BASICS, MAP SCALES & SYNOPTIC MAPS", "KPPSC PMS Geography Paper-I — Q.19: Cartography Basics: Map Scales & Synoptic Maps"),
        20: ("Q20_Map_Projections_Classification_and_Mercators_Projection_Master_Notes", "QUESTION 20: MAP PROJECTIONS: CLASSIFICATION & MERCATOR", "KPPSC PMS Geography Paper-I — Q.20: Map Projections: Classification & Mercator")
    }

    generated_count = 0
    for script_path, q_nums in script_map:
        print(f"\n=======================================================")
        print(f"Loading data from: {os.path.basename(script_path)}")
        print(f"=======================================================")
        ns = extract_data_from_file(script_path)
        
        for q_num in q_nums:
            meta = ns.get(f"p1_q{q_num}_meta")
            callout = ns.get(f"p1_q{q_num}_callout")
            cmap = ns.get(f"p1_q{q_num}_map")
            timeline = ns.get(f"p1_q{q_num}_timeline")
            essay = ns.get(f"p1_q{q_num}_essay")
            quotes = ns.get(f"p1_q{q_num}_quotes")
            
            if not all([meta, callout, cmap, timeline, essay, quotes]):
                print(f"WARNING: Incomplete data for Q{q_num} in {script_path}")
                continue
                
            # Build unified executive examination metadata
            new_meta = [
                ("Target Examination", "KPPSC PMS (Provincial Management Service) — Khyber Pakhtunkhwa"),
                ("Paper & Subject", "Geography Paper-I (Physical & General Geography — 100 Marks)"),
                ("EXAM QUESTIONS LINKED:", GEO_P1_EXAM_QUESTIONS_LINKED_MAP[q_num])
            ]
            for k, v in meta:
                if k not in ["Exam Years Solved", "Target Examination", "Paper & Subject"] and "EXAM QUESTIONS" not in k.upper():
                    new_meta.append((k, v))
                    
            fn_base, sub_t, full_t = titles_map[q_num]
            docx_file = os.path.join(GEO_P1_DIR, f"{fn_base}.docx")
            pdf_file = os.path.join(GEO_P1_DIR, f"{fn_base}.pdf")
            md_file = os.path.join(GEO_P1_DIR, f"{fn_base}.md")
            
            print(f"Generating Q{q_num}: {fn_base}...")
            generate_geo_master_doc(
                meta_data=new_meta,
                callout_data=callout,
                map_text=cmap,
                timeline_data=timeline,
                essay_sections=essay,
                quotes_data=quotes,
                out_path=docx_file,
                paper_sub_title=sub_t
            )
            
            generate_pdf_and_md(
                meta_data=new_meta,
                callout_data=callout,
                map_text=cmap,
                timeline_data=timeline,
                essay_sections=essay,
                quotes_data=quotes,
                pdf_path=pdf_file,
                md_path=md_file,
                title=full_t
            )
            generated_count += 1

    print(f"\nSuccessfully generated {generated_count} of 20 Geography Paper-I Master Topics!")

if __name__ == "__main__":
    update_and_build_all()
