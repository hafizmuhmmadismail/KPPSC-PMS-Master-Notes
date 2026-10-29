import os
import subprocess
import tempfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
PASHTO_DIR = r"H:\My Drive\PMS KPK 2022 notes\4. Pashto"

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_paragraph_rtl(p):
    pPr = p._element.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)

def create_docx(md_path, docx_path, title):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    in_code = False
    code_lines = []
    in_table = False
    table_rows = []
    
    for line in lines:
        raw_line = line.rstrip("\r\n")
        stripped = raw_line.strip()
        
        # Handle code blocks
        if stripped.startswith("```"):
            if in_code:
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                cell.width = Inches(6.8)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                code_lines = []
                in_code = False
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            else:
                in_code = True
                code_lines = []
            continue
            
        if in_code:
            code_lines.append(raw_line)
            continue
            
        # Handle table rows
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                table_rows = []
            if ":---" in stripped or "---:" in stripped or "| --- |" in stripped:
                continue
            cols = [c.strip() for c in stripped.split("|")[1:-1]]
            table_rows.append(cols)
            continue
        else:
            if in_table:
                num_rows = len(table_rows)
                num_cols = len(table_rows[0]) if num_rows > 0 else 0
                if num_rows > 0 and num_cols > 0:
                    tbl = doc.add_table(rows=num_rows, cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = False
                    for r_idx, row in enumerate(table_rows):
                        for c_idx, val in enumerate(row):
                            if c_idx < num_cols:
                                cell = tbl.cell(r_idx, c_idx)
                                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                                if r_idx == 0:
                                    set_cell_background(cell, "0E2B50")
                                elif r_idx % 2 == 1:
                                    set_cell_background(cell, "F8FAFC")
                                else:
                                    set_cell_background(cell, "FFFFFF")
                                    
                                p = cell.paragraphs[0]
                                is_ar = any('\u0600' <= char <= '\u06FF' for char in val)
                                if is_ar:
                                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                                    set_paragraph_rtl(p)
                                else:
                                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                    
                                text_clean = val.replace("<br>", "\n").replace("**", "")
                                run = p.add_run(text_clean)
                                run.font.name = "Bahij Nazanin" if is_ar else "Calibri"
                                run.font.size = Pt(9.5)
                                if r_idx == 0:
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                                else:
                                    run.font.color.rgb = RGBColor(30, 41, 59)
                    doc.add_paragraph().paragraph_format.space_after = Pt(4)
                in_table = False
                table_rows = []
                
        if not stripped:
            continue
            
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[2:])
            run.font.name = "Bahij Titr"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(14, 43, 80)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_paragraph_rtl(p)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[3:])
            run.font.name = "Bahij Titr"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(27, 73, 101)
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            is_ar = any('\u0600' <= char <= '\u06FF' for char in stripped)
            if is_ar:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_paragraph_rtl(p)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[4:])
            run.font.name = "Bahij Nazanin" if is_ar else "Calibri"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(14, 43, 80)
        elif stripped.startswith("> "):
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            cell.width = Inches(6.8)
            set_cell_background(cell, "F0F4F8")
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            p = cell.paragraphs[0]
            is_ar = any('\u0600' <= char <= '\u06FF' for char in stripped)
            if is_ar:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_paragraph_rtl(p)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[2:])
            run.font.name = "Bahij Zar" if is_ar else "Georgia"
            run.font.size = Pt(10.5)
            run.font.italic = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif stripped.startswith("---"):
            continue
        else:
            p = doc.add_paragraph()
            is_arabic = any('\u0600' <= char <= '\u06FF' for char in stripped)
            if is_arabic:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_paragraph_rtl(p)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(4)
            
            parts = stripped.split("**")
            for idx, part in enumerate(parts):
                if not part: continue
                run = p.add_run(part)
                run.font.name = "Bahij Nazanin" if is_arabic else "Calibri"
                run.font.size = Pt(11)
                if idx % 2 == 1:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(14, 43, 80)
                else:
                    run.font.color.rgb = RGBColor(40, 40, 40)
                    
    doc.save(docx_path)
    print(f"Created DOCX: {docx_path}")

def render_md_to_pdf(md_path, pdf_path, title):
    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    html_body = ""
    lines = md_text.split("\n")
    in_code = False
    code_block = []
    in_table = False
    table_rows = []
    
    for line in lines:
        raw_l = line.rstrip()
        s = raw_l.strip()
        
        if s.startswith("```"):
            if in_code:
                html_body += f"<pre style='direction:ltr; text-align:left; background:#f8fafc; border:1px solid #cbd5e1; border-radius:6px; padding:10px; font-family:Consolas, monospace; font-size:8.5pt; overflow-x:auto; line-height:1.3;'>{chr(10).join(code_block)}</pre>"
                code_block = []
                in_code = False
            else:
                in_code = True
                code_block = []
            continue
            
        if in_code:
            esc = raw_l.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            code_block.append(esc)
            continue
            
        if s.startswith("|") and s.endswith("|"):
            if not in_table:
                in_table = True
                table_rows = []
            if ":---" in s or "---:" in s or "| --- |" in s:
                continue
            cols = [c.strip() for c in s.split("|")[1:-1]]
            table_rows.append(cols)
            continue
        else:
            if in_table:
                t_html = "<table style='width:100%; border-collapse:collapse; margin:12px 0; font-size:10pt; direction:rtl;'>"
                for r_idx, row in enumerate(table_rows):
                    t_html += "<tr>"
                    tag = "th" if r_idx == 0 else "td"
                    bg = "#0E2B50" if r_idx == 0 else ("#f8fafc" if r_idx % 2 == 1 else "#ffffff")
                    color = "#ffffff" if r_idx == 0 else "#1e293b"
                    border = "1px solid #cbd5e1"
                    for col in row:
                        t_html += f"<{tag} style='padding:6px 10px; border:{border}; background:{bg}; color:{color}; font-family:\"Bahij Nazanin\", \"Segoe UI\", Tahoma;'>{col}</{tag}>"
                    t_html += "</tr>"
                t_html += "</table>"
                html_body += t_html
                in_table = False
                table_rows = []
                
        if not s:
            continue
            
        if s.startswith("# "):
            html_body += f"<h1 style='text-align:center; color:#0E2B50; font-family:\"Bahij Titr\", \"Segoe UI\", Tahoma; font-size:18pt; margin:14px 0 6px 0; border-bottom:2px solid #0E2B50; padding-bottom:6px;'>{s[2:]}</h1>"
        elif s.startswith("## "):
            html_body += f"<h2 style='text-align:right; direction:rtl; color:#1B4965; font-family:\"Bahij Titr\", \"Segoe UI\", Tahoma; font-size:14pt; margin:12px 0 4px 0;'>{s[3:]}</h2>"
        elif s.startswith("### "):
            is_ar = any('\u0600' <= char <= '\u06FF' for char in s)
            direction = "rtl" if is_ar else "ltr"
            align = "right" if is_ar else "left"
            html_body += f"<h3 style='text-align:{align}; direction:{direction}; color:#0E2B50; font-family:\"Bahij Nazanin\", \"Segoe UI\", Tahoma; font-size:12pt; margin:10px 0 2px 0; border-bottom:1px solid #e2e8f0; padding-bottom:2px;'>{s[4:]}</h3>"
        elif s.startswith("> "):
            is_ar = any('\u0600' <= char <= '\u06FF' for char in s)
            direction = "rtl" if is_ar else "ltr"
            align = "right" if is_ar else "left"
            border_side = "border-right:4px solid #1B4965; border-left:none;" if is_ar else "border-left:4px solid #1B4965; border-right:none;"
            html_body += f"<blockquote style='direction:{direction}; text-align:{align}; background:#F0F4F8; {border_side} padding:8px 14px; margin:10px 0; font-family:\"Bahij Zar\", Georgia, serif; font-size:10.5pt; color:#1e293b;'>{s[2:]}</blockquote>"
        elif s.startswith("---"):
            html_body += "<hr style='border:0; border-top:1px solid #cbd5e1; margin:14px 0;'>"
        else:
            is_arabic = any('\u0600' <= char <= '\u06FF' for char in s)
            direction = "rtl" if is_arabic else "ltr"
            align = "right" if is_arabic else "justify"
            font_family = "\"Bahij Nazanin\", \"Segoe UI\", Tahoma, sans-serif" if is_arabic else "Calibri, Arial, sans-serif"
            
            parts = s.split("**")
            formatted = ""
            for idx, part in enumerate(parts):
                if idx % 2 == 1:
                    formatted += f"<strong style='color:#0E2B50;'>{part}</strong>"
                else:
                    formatted += part
            html_body += f"<p style='direction:{direction}; text-align:{align}; font-family:{font_family}; font-size:11pt; line-height:1.6; margin:4px 0 8px 0;'>{formatted}</p>"

    full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset='utf-8'>
<style>
  @page {{ size: A4; margin: 18mm 16mm 18mm 16mm; }}
  body {{ font-family: 'Bahij Nazanin', 'Segoe UI', Tahoma, sans-serif; color: #1e293b; }}
  table {{ direction: rtl; }}
</style>
</head>
<body>
{html_body}
<div style='margin-top:20px; border-top:1px solid #94a3b8; padding-top:6px; text-align:center; font-size:9pt; color:#64748b; font-family:Calibri, sans-serif;'>
  KPPSC PMS Master Examination Notes Suite — Subject: Pashto (Optional) | Institutional 5-Page Standard
</div>
</body>
</html>"""

    with tempfile.NamedTemporaryFile("w", delete=False, suffix=".html", encoding="utf-8") as f:
        f.write(full_html)
        temp_html = f.name

    cmd = [
        EDGE_PATH,
        "--headless",
        "--disable-gpu",
        "--run-all-compositor-stages-before-draw",
        f"--print-to-pdf={pdf_path}",
        temp_html
    ]
    subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    if os.path.exists(temp_html):
        try: os.remove(temp_html)
        except: pass
    print(f"Created PDF: {pdf_path}")

def process_all_pashto_notes():
    files = [
        ("Q1_and_Q2_Pashto_Compulsory_Translations_and_Proverbs_2025_Master_Notes", "KPPSC PMS 2025 Pashto Q1 & Q2"),
        ("Q3_Abdul_Hamid_Baba_Poetic_Art_and_Imagery_2025_Master_Notes", "KPPSC PMS 2025 Pashto Q3 Abdul Hamid Baba"),
        ("Q4_Pashto_Folk_Literature_Folklore_Genres_2025_Master_Notes", "KPPSC PMS 2025 Pashto Q4 Folk Literature"),
        ("Q5_Evolution_of_Pashto_Drama_Radio_Stage_TV_2025_Master_Notes", "KPPSC PMS 2025 Pashto Q5 Pashto Drama"),
        ("Q6_Rahman_Baba_Poetic_Art_Aesthetics_and_Sufism_2025_Master_Notes", "KPPSC PMS 2025 Pashto Q6 Rahman Baba"),
        ("Q7_Modern_Pashto_Poets_Sayel_Salim_Raz_Azam_2025_Master_Notes", "KPPSC PMS 2025 Pashto Q7 Modern Poets"),
    ]
    
    for base_name, title in files:
        md_file = os.path.join(PASHTO_DIR, f"{base_name}.md")
        docx_file = os.path.join(PASHTO_DIR, f"{base_name}.docx")
        pdf_file = os.path.join(PASHTO_DIR, f"{base_name}.pdf")
        
        if os.path.exists(md_file):
            print(f"Processing: {base_name}...")
            create_docx(md_file, docx_file, title)
            render_md_to_pdf(md_file, pdf_file, title)

if __name__ == "__main__":
    process_all_pashto_notes()
